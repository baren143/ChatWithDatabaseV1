"""File upload + document management endpoints.

Auth: every endpoint requires a valid Bearer token. The hardcoded
`test_user_123` user is gone.

Storage: each uploaded file is stored under
`uploads/<user_id>/<uuid>__<safe_filename>`. This prevents one user
from clobbering another user's file when they upload files with the
same name, and protects the original filename for display purposes
while the storage path is content-independent.
"""

from __future__ import annotations

import logging
import os
import re
import uuid
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from fastapi.responses import FileResponse, Response
from sqlalchemy import and_
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool

from auth.utils import get_current_user
from database import get_db
from dependencies import resolve_user_id_from_request
from models import Document, DocumentRow, DocumentVector, User
from tasks import process_document

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api", tags=["upload"])

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB
ALLOWED_EXTENSIONS = {
    ".pdf", ".txt", ".csv", ".xlsx", ".xls", ".doc", ".docx", ".md",
}

# Per-extension magic-byte / parser validation. Prevents uploading
# `malware.exe` renamed to `malware.xlsx` and having the processor
# try to parse random binary as a spreadsheet.
_VALIDATORS = {
    ".xlsx": lambda p: _validate_xlsx(p),
    ".xls": lambda p: _validate_xls(p),
    ".csv": lambda p: _validate_csv(p),
    ".pdf": lambda p: _validate_pdf(p),
    ".txt": lambda p: _validate_text(p),
    ".md": lambda p: _validate_text(p),
    ".docx": lambda p: _validate_docx(p),
    ".doc": lambda p: _validate_doc(p),
}

# Per-extension MIME types we record on the Document row.
_MIME_TYPES = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
}

# Uploads root — derived from this file's location, not cwd.
UPLOADS_ROOT = Path(__file__).resolve().parent.parent / "uploads"


# ── Helpers ──────────────────────────────────────────────────────────────


def _sanitize_filename(filename: str) -> str:
    """Strip path components and any control / shell metacharacters."""
    filename = Path(filename).name
    # Allow letters, digits, dot, dash, underscore, space, parens.
    filename = re.sub(r"[^\w\-. ()]", "_", filename)
    # Collapse runs of underscores that may result from sanitization.
    filename = re.sub(r"_+", "_", filename).strip(" _.")
    if not filename:
        filename = f"upload_{uuid.uuid4().hex[:8]}"
    if len(filename) > 200:
        stem, _, ext = filename.rpartition(".")
        ext = ("." + ext) if ext else ""
        keep = 200 - len(ext)
        filename = stem[: max(1, keep)] + ext
    return filename


def _is_allowed(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def _user_upload_dir(user_id: str) -> Path:
    """Per-user upload directory. Auto-created on first write."""
    d = UPLOADS_ROOT / user_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def _storage_path(user_id: str, safe_filename: str) -> tuple[str, str]:
    """Return (storage_key, full_path). The storage_key is stored in the DB
    so renames in the UI don't move the actual file."""
    storage_key = f"{uuid.uuid4().hex}__{safe_filename}"
    full_path = _user_upload_dir(user_id) / storage_key
    return storage_key, str(full_path)


def _is_image(blob: bytes) -> bool:
    """Check if the blob is an image by magic bytes."""
    # Check for common image signatures
    if blob.startswith(b'\x89PNG\r\n\x1a\n'):  # PNG
        return True
    if blob.startswith(b'\xff\xd8\xff'):  # JPEG
        return True
    if blob.startswith(b'GIF87a') or blob.startswith(b'GIF89a'):  # GIF
        return True
    if blob.startswith(b'BM'):  # BMP
        return True
    if blob.startswith(b'II*\x00') or blob.startswith(b'MM\x00*'):  # TIFF
        return True
    # WebP
    if blob.startswith(b'RIFF') and blob[8:12] == b'WEBP':
        return True
    return False


# ── Content validators (each raises HTTPException 400 on bad payload) ───────

def _validate_xlsx(path: Path) -> None:
    try:
        import openpyxl
        with open(path, "rb") as f:
            openpyxl.load_workbook(f, read_only=True, data_only=True).close()
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File is not a valid XLSX: {e}",
        ) from e


def _validate_xls(path: Path) -> None:
    try:
        import xlrd  # type: ignore
        with open(path, "rb") as f:
            xlrd.open_workbook(file_contents=f.read())
    except ImportError:
        # If xlrd isn't installed, we still want to accept the file and
        # surface the error during processing instead of upload.
        return
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File is not a valid XLS: {e}",
        ) from e


def _validate_csv(path: Path) -> None:
    try:
        import pandas as pd
        pd.read_csv(path, nrows=5)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File is not a valid CSV: {e}",
        ) from e


def _validate_pdf(path: Path) -> None:
    try:
        from pypdf import PdfReader
        with open(path, "rb") as f:
            PdfReader(f)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File is not a valid PDF: {e}",
        ) from e


def _validate_text(path: Path) -> None:
    try:
        with open(path, "r", encoding="utf-8") as f:
            f.read(1024)
    except UnicodeDecodeError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File is not valid UTF-8 text.",
        )


def _validate_docx(path: Path) -> None:
    try:
        import docx  # type: ignore
        with open(path, "rb") as f:
            docx.Document(f)
    except ImportError:
        return
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File is not a valid DOCX: {e}",
        ) from e


def _validate_doc(path: Path) -> None:
    # We don't have a reliable parser for legacy .doc; rely on extension +
    # processing-time handling.
    return


# ── Endpoints ────────────────────────────────────────────────────────────

def _process_upload_file(
    tmp_path: Path,
    full_path: str,
    blob: bytes,
    ext: str,
    safe_filename: str,
    storage_key: str,
    user_id: str,
    db: Session,
    validator,
    mime_type: Optional[str],
) -> Document:
    # 1. Write the blob to tmp_path
    try:
        with open(tmp_path, "wb") as f:
            f.write(blob)
    except Exception as e:
        logger.exception("Failed to write upload buffer for user %s", user_id)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to save file: {e}",
        ) from e

    # 2. Content validation (real parser, not just extension)
    if validator is not None:
        try:
            validator(tmp_path)
        except HTTPException:
            tmp_path.unlink(missing_ok=True)
            raise
        except Exception as e:
            tmp_path.unlink(missing_ok=True)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"File failed validation: {e}",
            ) from e

    # 3. Atomic rename
    try:
        os.replace(tmp_path, full_path)
    except Exception as e:
        tmp_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to finalize upload: {e}",
        ) from e

    # 4. Persist the Document row.
    try:
        document = Document(
            user_id=user_id,
            file_name=safe_filename,
            file_path=full_path,
            storage_key=storage_key,
            file_size=len(blob),
            mime_type=mime_type,
            status="processing",
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        return document
    except Exception as e:
        # cleanup full_path if DB write fails
        try:
            if os.path.exists(full_path):
                os.remove(full_path)
        except Exception:
            pass
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Database error during upload: {e}",
        ) from e


@router.post("/upload")
async def upload_file(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    user_id = resolve_user_id_from_request(request, db)

    if not file.filename or not _is_allowed(file.filename):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "File type not allowed. Supported: "
                + ", ".join(sorted(ALLOWED_EXTENSIONS))
            ),
        )

    # Read the upload into a buffer so we can validate size + content
    # atomically without leaving a partial file on disk.
    blob = await file.read()
    if len(blob) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds the {MAX_FILE_SIZE // (1024 * 1024)} MB limit.",
        )
    if len(blob) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty.",
        )

    # Check if the file is an image (by magic bytes) and reject if so
    if _is_image(blob):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Image files are not supported.",
        )

    safe_filename = _sanitize_filename(file.filename)
    ext = Path(safe_filename).suffix.lower()
    storage_key, full_path = _storage_path(user_id, safe_filename)

    # Write to a tmp path first so a failed validation never leaves a
    # half-written real file on disk.
    tmp_path = Path(full_path + ".part")

    # Content validator
    validator = _VALIDATORS.get(ext)
    mime_type = _MIME_TYPES.get(ext)

    # Offload the blocking operations to a worker thread
    document = await run_in_threadpool(
        _process_upload_file,
        tmp_path=tmp_path,
        full_path=full_path,
        blob=blob,
        ext=ext,
        safe_filename=safe_filename,
        storage_key=storage_key,
        user_id=user_id,
        db=db,
        validator=validator,
        mime_type=mime_type,
    )

    # Kick off async processing
    try:
        process_document.delay(document.id)
    except Exception as e:
        # Redis / Celery down — flip to error so the client sees it
        logger.exception("Failed to enqueue processing for %s", document.id)
        document.status = "error"
        document.error_message = "Background worker unavailable. Please retry shortly."
        db.commit()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Background worker unavailable. Try again in a moment.",
        ) from e

    return {
        "id": document.id,
        "document_id": document.id,
        "file_name": document.file_name,
        "status": document.status,
        "size": document.file_size,
    }


@router.get("/documents")
def list_documents(
    request: Request,
    db: Session = Depends(get_db),
    skip: int = 0,
    limit: int = 50,
):
    user_id = resolve_user_id_from_request(request, db)
    total = db.query(Document).filter(Document.user_id == user_id).count()
    docs = (
        db.query(Document)
        .filter(Document.user_id == user_id)
        .order_by(Document.created_at.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return {
        "total": total,
        "skip": skip,
        "limit": limit,
        "items": [
            {
                "id": d.id,
                "file_name": d.file_name,
                "status": d.status,
                "total_chunks": d.total_chunks or 0,
                "processed_chunks": d.processed_chunks or 0,
                "file_size": d.file_size or 0,
                "created_at": d.created_at.isoformat() if d.created_at else None,
            }
            for d in docs
        ],
    }


@router.get("/documents/{document_id}")
def get_document(
    document_id: str,
    request: Request,
    db: Session = Depends(get_db),
):
    user_id = resolve_user_id_from_request(request, db)
    doc = (
        db.query(Document)
        .filter(and_(Document.id == document_id, Document.user_id == user_id))
        .first()
    )
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found",
        )
    return {
        "id": doc.id,
        "file_name": doc.file_name,
        "status": doc.status,
        "total_chunks": doc.total_chunks or 0,
        "processed_chunks": doc.processed_chunks or 0,
        "file_size": doc.file_size or 0,
        "error_message": doc.error_message,
        "created_at": doc.created_at.isoformat() if doc.created_at else None,
    }


@router.delete("/documents/{document_id}")
def delete_document(
    document_id: str,
    request: Request,
    db: Session = Depends(get_db),
):
    user_id = resolve_user_id_from_request(request, db)
    doc = (
        db.query(Document)
        .filter(and_(Document.id == document_id, Document.user_id == user_id))
        .first()
    )
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found",
        )

    # Delete related rows + vectors
    db.query(DocumentRow).filter(DocumentRow.document_id == document_id).delete()
    db.query(DocumentVector).filter(DocumentVector.document_id == document_id).delete()
    # DocumentRow cascades from Document, but be explicit for safety
    db.query(Document).filter(Document.id == document_id).delete()
    db.commit()

    # Best-effort physical file removal
    try:
        if doc.file_path and os.path.exists(doc.file_path):
            os.remove(doc.file_path)
    except Exception:
        logger.exception("Failed to remove file for %s", document_id)

    return {"status": "success", "message": "Document deleted successfully"}